"""
Building definition (BUI) for pyBuildingEnergy / ISO 52016-1, read from Excel.

All inputs live in building_input.xlsx (sheets 'Building', 'Surfaces', 'Profiles').
This module reads the workbook and builds the BUI dictionary that run_demand.py uses:

    from my_building import BUI

To use another workbook, change EXCEL_FILE below, or call load_bui_from_excel(path).
"""
from pathlib import Path

from openpyxl import load_workbook

EXCEL_FILE = Path(__file__).resolve().parent / "building_input.xlsx"

# Layout of the workbook (must match the template)
HEADER_ROW = 4          # header row with the keys on 'Surfaces' and 'Profiles'
SURFACES_FIRST_ROW = 6  # first data row on 'Surfaces' (row 5 holds labels/units)
PROFILES_FIRST_ROW = 5  # hour 0 on 'Profiles'

GAIN_NAMES = ["occupants", "appliances", "lighting"]
HVAC_PROFILES = ["heating", "cooling", "ventilation"]


class ExcelInputError(ValueError):
    """Raised when the Excel input is missing data or has invalid values."""


# ----------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------
def _is_blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def _num(value, where, required=True, default=None):
    if _is_blank(value):
        if required:
            raise ExcelInputError(f"Missing value: {where}")
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        raise ExcelInputError(f"Not a number: {where} = {value!r}")


def _bool(value, where):
    if _is_blank(value):
        return False
    if isinstance(value, bool):
        return value
    s = str(value).strip().lower()
    if s in ("true", "yes", "1", "y"):
        return True
    if s in ("false", "no", "0", "n"):
        return False
    raise ExcelInputError(f"Expected TRUE/FALSE: {where} = {value!r}")


def _header_map(ws):
    """Map header key -> column index from the header row."""
    cols = {}
    for cell in ws[HEADER_ROW]:
        if not _is_blank(cell.value):
            cols[str(cell.value).strip()] = cell.column
    return cols


# ----------------------------------------------------------------------
# Sheet readers
# ----------------------------------------------------------------------
def _read_building_sheet(ws):
    """Read key/value pairs: key in column A, value in column C."""
    values = {}
    for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        key = row[0] if len(row) > 0 else None
        if _is_blank(key):
            continue
        values[str(key).strip()] = row[2] if len(row) > 2 else None
    return values


def _read_surfaces(ws):
    cols = _header_map(ws)
    for required in ["name", "type", "area", "u_value", "azimuth", "tilt"]:
        if required not in cols:
            raise ExcelInputError(f"'Surfaces' sheet: column '{required}' not found in row {HEADER_ROW}")

    def get(row, key):
        c = cols.get(key)
        return ws.cell(row=row, column=c).value if c else None

    surfaces = []
    row = SURFACES_FIRST_ROW
    while not _is_blank(get(row, "name")):
        where = lambda key: f"Surfaces row {row}, '{key}'"
        name = str(get(row, "name")).strip()
        s_type = str(get(row, "type") or "").strip().lower()
        if s_type not in ("opaque", "transparent"):
            raise ExcelInputError(f"{where('type')}: must be 'opaque' or 'transparent', got {get(row, 'type')!r}")
        boundary = str(get(row, "boundary") or "OUTDOORS").strip().upper()
        if boundary not in ("OUTDOORS", "GROUND"):
            raise ExcelInputError(f"{where('boundary')}: must be OUTDOORS or GROUND, got {boundary!r}")

        surf = {
            "name": name,
            "type": s_type,
            "area": _num(get(row, "area"), where("area")),
            "sky_view_factor": _num(get(row, "sky_view_factor"), where("sky_view_factor"), False, 0.5),
            "u_value": _num(get(row, "u_value"), where("u_value")),
            "solar_absorptance": _num(get(row, "solar_absorptance"), where("solar_absorptance"), False, 0.5),
            "thermal_capacity": _num(get(row, "thermal_capacity"), where("thermal_capacity"), False, 0.0),
            "orientation": {
                "azimuth": _num(get(row, "azimuth"), where("azimuth")),
                "tilt": _num(get(row, "tilt"), where("tilt")),
            },
            "name_adj_zone": None,
            "height": _num(get(row, "height"), where("height"), False, 0.0),
        }
        if boundary == "GROUND":
            surf["boundary"] = "GROUND"

        if s_type == "opaque":
            surf["length"] = _num(get(row, "length"), where("length"), False, 0.0)
        else:
            surf.update({
                "g_value": _num(get(row, "g_value"), where("g_value")),
                "width": _num(get(row, "width"), where("width"), False, 0.0),
                "parapet": _num(get(row, "parapet"), where("parapet"), False, 0.0),
                "shading": _bool(get(row, "shading"), where("shading")),
                "shading_type": (str(get(row, "shading_type")).strip()
                                 if not _is_blank(get(row, "shading_type")) else "horizontal_overhang"),
                "width_or_distance_of_shading_elements": _num(
                    get(row, "width_or_distance_of_shading_elements"),
                    where("width_or_distance_of_shading_elements"), False, 0.0),
                "overhang_properties": {
                    "width_of_horizontal_overhangs": _num(
                        get(row, "width_of_horizontal_overhangs"),
                        where("width_of_horizontal_overhangs"), False, 0.0),
                },
            })
        surfaces.append(surf)
        row += 1

    if not surfaces:
        raise ExcelInputError("'Surfaces' sheet: no surfaces found")
    return surfaces


def _read_profiles(ws):
    cols = _header_map(ws)
    profiles = {}
    for name in GAIN_NAMES + HVAC_PROFILES:
        for day in ("weekday", "weekend"):
            key = f"{name}_{day}"
            if key not in cols:
                raise ExcelInputError(f"'Profiles' sheet: column '{key}' not found in row {HEADER_ROW}")
            vals = []
            for h in range(24):
                r = PROFILES_FIRST_ROW + h
                v = _num(ws.cell(row=r, column=cols[key]).value, f"Profiles '{key}', hour {h}")
                if not 0.0 <= v <= 1.0:
                    raise ExcelInputError(f"Profiles '{key}', hour {h}: value {v} must be between 0 and 1")
                vals.append(v)
            profiles[key] = vals
    return profiles


# ----------------------------------------------------------------------
# Main loader
# ----------------------------------------------------------------------
def load_bui_from_excel(path=EXCEL_FILE):
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Excel input not found: {path}")
    wb = load_workbook(path, data_only=True)
    for sheet in ("Building", "Surfaces", "Profiles"):
        if sheet not in wb.sheetnames:
            raise ExcelInputError(f"Sheet '{sheet}' missing in {path.name}")

    g = _read_building_sheet(wb["Building"])
    n = lambda key, required=True, default=None: _num(g.get(key), f"Building '{key}'", required, default)
    t = lambda key, default="": str(g.get(key)).strip() if not _is_blank(g.get(key)) else default

    profiles = _read_profiles(wb["Profiles"])
    wall_thickness = n("wall_thickness", False, 0.3)

    return {
        "building": {
            "name": t("name", "Building"),
            "azimuth_relative_to_true_north": n("azimuth_relative_to_true_north", False, 0.0),
            "latitude": n("latitude"),
            "longitude": n("longitude"),
            "exposed_perimeter": n("exposed_perimeter"),
            "height": n("height"),
            "wall_thickness": wall_thickness,
            "n_floors": int(n("n_floors")),
            "building_type_class": t("building_type_class", "Residential_apartment"),
            "adj_zones_present": False,
            "number_adj_zone": 0,
            "net_floor_area": n("net_floor_area"),
            "construction_class": t("construction_class", "class_i"),
            "construction_year": t("construction_year"),
            "country": t("country"),
        },
        "adjacent_zones": [],
        "building_surface": _read_surfaces(wb["Surfaces"]),
        "units": {
            "area": "m²",
            "u_value": "W/m²K",
            "thermal_capacity": "J/m²K",
            "azimuth": "degrees (0=N, 90=E, 180=S, 270=W)",
            "tilt": "degrees (0=horizontal, 90=vertical)",
            "internal_gain": "W/m²",
            "internal_gain_profile": "Normalized to 0-1",
            "HVAC_profile": "0: off, 1: on",
        },
        "building_parameters": {
            "temperature_setpoints": {
                "heating_setpoint": n("heating_setpoint"),
                "heating_setback": n("heating_setback"),
                "cooling_setpoint": n("cooling_setpoint"),
                "cooling_setback": n("cooling_setback"),
                "units": "°C",
            },
            "system_capacities": {
                "heating_capacity": n("heating_capacity"),
                "cooling_capacity": n("cooling_capacity"),
                "units": "W",
            },
            "ventilation": {
                "ventilation_type": t("ventilation_type", "occupancy").lower(),
                "flow_rate_per_person": n("flow_rate_per_person", False, 0.0),
                "units": "l/(s m2)",
                "custom_heat_transfer_coefficient_ventilation": n(
                    "custom_heat_transfer_coefficient_ventilation", False, 0.0),
            },
            "internal_gains": [
                {
                    "name": name,
                    "full_load": n(f"gain_{name}"),
                    "weekday": profiles[f"{name}_weekday"],
                    "weekend": profiles[f"{name}_weekend"],
                }
                for name in GAIN_NAMES
            ],
            "construction": {
                "wall_thickness": wall_thickness,
                "thermal_bridge_heat_W_K": n("thermal_bridge_heat_W_K", False, 0.0),
                "units": "m (for thickness), W/K (for total thermal-bridge coefficient)",
            },
            "climate_parameters": {
                "coldest_month": int(n("coldest_month", False, 1)),
                "units": "1-12 (January-December)",
            },
            **{
                f"{k}_profile": {"weekday": profiles[f"{k}_weekday"], "weekend": profiles[f"{k}_weekend"]}
                for k in HVAC_PROFILES
            },
        },
    }


BUI = load_bui_from_excel(EXCEL_FILE)


if __name__ == "__main__":
    # Quick check: python my_building.py
    b = BUI["building"]
    print(f"Read '{b['name']}' from {EXCEL_FILE.name}: "
          f"{b['net_floor_area']} m², {len(BUI['building_surface'])} surfaces")
    for s in BUI["building_surface"]:
        print(f"  {s['name']:<28} {s['type']:<12} A={s['area']:>7.1f} m²  U={s['u_value']:.2f}  "
              f"az={s['orientation']['azimuth']:.0f}  tilt={s['orientation']['tilt']:.0f}  "
              f"{s.get('boundary', 'OUTDOORS')}")
